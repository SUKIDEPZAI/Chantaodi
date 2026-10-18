"""Xuất file (docx/xlsx) + giao thức tạo file/suy luận cho MotAPI Studio."""
import io, re, csv
from flask import Blueprint, request, send_file, jsonify

bp = Blueprint('export', __name__)

FILE_PROTOCOL = (
    "\n\nQUY ƯỚC ĐẦU RA:\n"
    "1) Câu hỏi phức tạp: mở đầu bằng một khối <think>…</think> gồm 3-6 gạch đầu dòng ngắn tóm tắt cách tiếp cận, giả định, điểm cần kiểm chứng; rồi mới trả lời. Câu hỏi đơn giản thì bỏ qua.\n"
    "2) Khi người dùng yêu cầu tạo/xuất file (docx, xlsx, csv, md, txt, html, json, py, ...), đặt TOÀN BỘ nội dung file trong một khối "
    "```file:tên-file.đuôi ... ``` rồi viết 1 câu mô tả. docx/pdf: nội dung viết bằng Markdown; xlsx: nội dung dạng CSV.\n"
    "3) Khi viết giao diện web/SVG hoàn chỉnh: gộp thành MỘT tệp ```html tự chứa (CSS/JS nội tuyến) để người dùng xem trước ngay.\n"
)


def _clean(t):
    return re.sub(r'\*\*|`|__', '', t)


@bp.post('/api/export')
def export():
    d = request.get_json(silent=True) or {}
    fmt = str(d.get('format', 'docx')).lower()
    title = str(d.get('title') or 'document')[:80]
    content = str(d.get('content') or '')[:400000]
    safe = re.sub(r'[^\w\-. ]+', '_', title).strip() or 'document'
    try:
        if fmt == 'docx':
            from docx import Document
            from docx.shared import Pt
            doc = Document(); code = False
            for line in content.split('\n'):
                if line.strip().startswith('```'):
                    code = not code; continue
                if code:
                    r = doc.add_paragraph().add_run(line); r.font.name = 'Consolas'; r.font.size = Pt(9); continue
                m = re.match(r'^(#{1,3}) (.+)', line)
                if m: doc.add_heading(_clean(m.group(2)), len(m.group(1))); continue
                if re.match(r'^\s*[-*] ', line): doc.add_paragraph(_clean(re.sub(r'^\s*[-*] ', '', line)), style='List Bullet'); continue
                if re.match(r'^\s*\d+\. ', line): doc.add_paragraph(_clean(re.sub(r'^\s*\d+\. ', '', line)), style='List Number'); continue
                if line.strip(): doc.add_paragraph(_clean(line))
            buf = io.BytesIO(); doc.save(buf); buf.seek(0)
            return send_file(buf, as_attachment=True, download_name=safe + '.docx',
                             mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        if fmt == 'xlsx':
            from openpyxl import Workbook
            wb = Workbook(); ws = wb.active
            for row in csv.reader(io.StringIO(content)): ws.append(row)
            buf = io.BytesIO(); wb.save(buf); buf.seek(0)
            return send_file(buf, as_attachment=True, download_name=safe + '.xlsx',
                             mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        if fmt == 'pptx':
            from pptx import Presentation
            from pptx.util import Pt
            prs = Presentation(); slides = []; cur = None
            for line in content.split('\n'):
                m = re.match(r'^(#{1,2}) (.+)', line)
                if m: cur = [_clean(m.group(2)), []]; slides.append(cur); continue
                if line.strip() and not line.strip().startswith('```'):
                    if cur is None: cur = [title, []]; slides.append(cur)
                    cur[1].append(_clean(re.sub(r'^\s*([-*]|\d+\.) ', '', line)))
            for t, b in (slides or [[title, [content[:500]]]]):
                s = prs.slides.add_slide(prs.slide_layouts[1]); s.shapes.title.text = t[:120]
                tf = s.placeholders[1].text_frame; tf.text = '\n'.join(b[:10])
                for p in tf.paragraphs: p.font.size = Pt(20)
            buf = io.BytesIO(); prs.save(buf); buf.seek(0)
            return send_file(buf, as_attachment=True, download_name=safe + '.pptx',
                             mimetype='application/vnd.openxmlformats-officedocument.presentationml.presentation')
        if fmt == 'rtf':
            def esc(t):
                return ''.join('\\' + c if c in '\\{}' else (c if ord(c) < 128 else '\\u%d?' % (ord(c) if ord(c) < 32768 else ord(c) - 65536)) for c in t)
            body = '\\par\n'.join(esc(_clean(l)) for l in content.split('\n'))
            data = '{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 Arial;}}\\f0\\fs24 ' + body + '}'
            return send_file(io.BytesIO(data.encode('ascii')), as_attachment=True, download_name=safe + '.rtf', mimetype='application/rtf')
        return jsonify(error='Định dạng chưa hỗ trợ trên máy chủ.'), 400
    except ImportError:
        return jsonify(error='Thiếu thư viện: kiểm tra python-docx, openpyxl, python-pptx trong requirements.txt.'), 501
    except Exception:
        return jsonify(error='Không tạo được tệp. Kiểm tra nội dung và log Render.'), 500
